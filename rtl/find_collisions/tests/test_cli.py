from __future__ import annotations

import gzip
from pathlib import Path

import pytest
from conftest import make_cfg_sv, make_fullchipdump, make_xml
from openpyxl import load_workbook

from find_collisions.cli import main

PARENT = "parent cell's library"
SEARCH = "default library search order"


def rewrite_dump(workarea: Path, instances: list[tuple[str, str, str, str]]) -> None:
    xml = workarea.parent / "config_diagnostics.xml.gz"
    xml.write_bytes(gzip.compress(make_xml(instances).encode()))


def run(argv: list[str], cwd: Path, monkeypatch: pytest.MonkeyPatch) -> int:
    monkeypatch.chdir(cwd)
    return main(argv)


def test_missing_fe_collateral_exits_two(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    assert run(["--fe_collateral", str(tmp_path / "nope")], tmp_path, monkeypatch) == 2


def test_dry_run_writes_nothing(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture
) -> None:
    outdir = tmp_path / "out"
    outdir.mkdir()
    assert run(["--fe_collateral", str(workarea), "--dry-run"], outdir, monkeypatch) == 0
    assert list(outdir.iterdir()) == []
    assert "would write" in capsys.readouterr().out


def test_run_writes_reports_in_cwd(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    rewrite_dump(
        workarea,
        [
            ("corimh", "IMH", "Top Module", '"/rtl/corimh.sv",1'),
            ("alpha", "LIB_A", PARENT, '"/rtl/a/alpha.sv",10'),
            ("alpha", "LIB_A", PARENT, '"/rtl/a/alpha.sv",10'),
            ("alpha", "LIB_B", SEARCH, '"/rtl/b/alpha.sv",20'),
            ("beta", "LIB_A", PARENT, '"/rtl/a/beta.sv",5'),
        ],
    )
    outdir = tmp_path / "out"
    outdir.mkdir()
    assert run(["--fe_collateral", str(workarea), "--verbose"], outdir, monkeypatch) == 0

    workbook = load_workbook(outdir / "find_collisions.xlsx")
    assert workbook.sheetnames == ["raw", "reduced"]
    raw = [[c.value for c in row] for row in workbook["raw"].iter_rows(min_row=2)]
    assert [(r[0], r[1], r[4]) for r in raw] == [
        ("alpha", "LIB_A", 2),
        ("alpha", "LIB_B", 1),
    ]
    reduced = [[c.value for c in row] for row in workbook["reduced"].iter_rows(min_row=2)]
    assert [(r[0], r[1], r[4]) for r in reduced] == [
        ("alpha", '"/rtl/a/alpha.sv",10', 2),
        ("alpha", '"/rtl/b/alpha.sv",20', 1),
    ]
    header = [c.value for c in workbook["reduced"][1]]
    assert set(header[5:]) == {"Top Module", PARENT, SEARCH}

    report = (outdir / "find_collisions.report").read_text()
    assert "    raw: 2" in report
    assert "    reduced: 2" in report
    assert not (outdir / "find_collisions.module.xlsx").exists()


def test_relative_path_false_positive_dropped_from_reduced(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    real = tmp_path / "shared.sv"
    real.write_text("")
    alias_a = f'"{tmp_path / "a" / ".." / "shared.sv"}",7'
    alias_b = f'"{tmp_path / "b" / ".." / "shared.sv"}",7'
    rewrite_dump(
        workarea,
        [
            ("corimh", "IMH", "Top Module", '"/rtl/corimh.sv",1'),
            ("arf_clk", "LIB_A", PARENT, alias_a),
            ("arf_clk", "LIB_B", SEARCH, alias_b),
        ],
    )
    outdir = tmp_path / "out"
    outdir.mkdir()
    assert run(["--fe_collateral", str(workarea)], outdir, monkeypatch) == 0

    workbook = load_workbook(outdir / "find_collisions.xlsx")
    assert workbook["raw"].max_row == 3
    assert workbook["reduced"].max_row == 1


def test_module_selects_module_workbook(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    rewrite_dump(
        workarea,
        [
            ("corimh", "IMH", "Top Module", '"/rtl/corimh.sv",1'),
            ("alpha", "LIB_A", PARENT, '"/rtl/a/alpha.sv",10'),
            ("alpha", "LIB_B", SEARCH, '"/rtl/b/alpha.sv",20'),
            ("gamma", "LIB_A", PARENT, '"/rtl/a/gamma.sv",1'),
            ("gamma", "LIB_B", PARENT, '"/rtl/b/gamma.sv",1'),
        ],
    )
    outdir = tmp_path / "out"
    outdir.mkdir()
    assert run(["--fe_collateral", str(workarea), "--module", "gamma"], outdir, monkeypatch) == 0

    raw = load_workbook(outdir / "find_collisions.xlsx")["raw"]
    assert {row[0].value for row in raw.iter_rows(min_row=2)} == {"alpha", "gamma"}

    module_book = load_workbook(outdir / "find_collisions.module.xlsx")
    assert {row[0].value for row in module_book["raw"].iter_rows(min_row=2)} == {"gamma"}
    assert list(module_book["patterns"].iter_rows(min_row=2, values_only=True)) == [
        ("gamma", "exact", 1)
    ]
    report = (outdir / "find_collisions.report").read_text()
    assert "patterns:\n  gamma" in report


def test_module_and_modulefile_are_exclusive(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    argv = ["--fe_collateral", str(workarea), "--module", "a", "--modulefile", "b"]
    with pytest.raises(SystemExit):
        run(argv, tmp_path, monkeypatch)


def test_paranoia_reports_and_never_fails(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    instances = [
        ("corimh", "IMH", "Top Module", '"/rtl/corimh.sv",1'),
        ("alpha", "LIB_A", PARENT, '"/rtl/a/alpha.sv",10'),
        ("beta", "LIB_B", SEARCH, '"/rtl/b/beta.sv",20'),
    ]
    rewrite_dump(workarea, instances)
    fullchipdump = workarea.parent / "fev_v2k" / "fullchipdump.final.py.sort.zst"
    fullchipdump.write_text(make_fullchipdump(instances[:2]))
    (workarea / "corimh_cfg.sv").write_text(make_cfg_sv(instances))

    outdir = tmp_path / "out"
    outdir.mkdir()
    assert run(["--fe_collateral", str(workarea), "--paranoia"], outdir, monkeypatch) == 0

    report = (outdir / "find_collisions.report").read_text()
    assert "paranoia results:" in report
    assert "  xml instance count:          3" in report
    assert "  fullchipdump instance count: 2" in report
    assert "  differences found: 1" in report
    assert "    xml instances missing from the fullchipdump: 1" in report
    assert "      { 'instance' : r\"corimh.beta_inst\"" in report
    assert "    instances in corimh_cfg.sv: 2" in report
    assert "    v2k config binding mismatch from xml: 0" in report


MODULE_DUMP = [
    ("corimh", "IMH", "Top Module", '"/rtl/corimh.sv",1'),
    ("xor_gate", "LIB_A", PARENT, '"/rtl/a/xor_gate.sv",10'),
    ("xor_gate", "LIB_B", SEARCH, '"/rtl/b/xor_gate.sv",20'),
    ("fblp_fcr_ctrl", "LIB_A", PARENT, '"/rtl/a/fblp_fcr_ctrl.sv",5'),
    ("my_xor_gate", "LIB_A", PARENT, '"/rtl/a/my_xor_gate.sv",5'),
    ("and_gate", "LIB_A", PARENT, '"/rtl/a/and_gate.sv",1'),
    ("and_gate", "LIB_B", SEARCH, '"/rtl/b/and_gate.sv",2'),
]


def test_modulefile_writes_module_workbook(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    rewrite_dump(workarea, MODULE_DUMP)
    modulefile = tmp_path / "modules.md"
    modulefile.write_text('# Test module list\nxor_gate\nr"fblp_fcr"\n')
    outdir = tmp_path / "out"
    outdir.mkdir()
    assert (
        run(
            ["--fe_collateral", str(workarea), "--modulefile", str(modulefile)],
            outdir,
            monkeypatch,
        )
        == 0
    )

    workbook = load_workbook(outdir / "find_collisions.module.xlsx")
    assert workbook.sheetnames == ["patterns", "raw", "reduced", "instances"]

    patterns = [row for row in workbook["patterns"].iter_rows(min_row=2, values_only=True)]
    assert patterns == [("xor_gate", "exact", 1), ('r"fblp_fcr"', "regex", 1)]

    raw = [row[0] for row in workbook["raw"].iter_rows(min_row=2, values_only=True)]
    assert set(raw) == {"xor_gate"}  # fblp_fcr_ctrl has a single definition

    reduced = [row[0] for row in workbook["reduced"].iter_rows(min_row=2, values_only=True)]
    assert set(reduced) == {"xor_gate"}

    instances = list(workbook["instances"].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in instances] == ["fblp_fcr_ctrl", "xor_gate", "xor_gate"]
    assert all(row[4].startswith("corimh.") for row in instances)
    assert all(row[5].startswith('"') for row in instances)

    report = (outdir / "find_collisions.report").read_text()
    assert str(modulefile) in report
    assert "    instances: 3" in report


def test_modulefile_regex_matches_substring(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    rewrite_dump(workarea, MODULE_DUMP)
    modulefile = tmp_path / "modules.md"
    modulefile.write_text('r"xor"\n')
    outdir = tmp_path / "out"
    outdir.mkdir()
    run(["--fe_collateral", str(workarea), "--modulefile", str(modulefile)], outdir, monkeypatch)

    workbook = load_workbook(outdir / "find_collisions.module.xlsx")
    assert workbook["patterns"]["C2"].value == 2
    modules = {row[0] for row in workbook["instances"].iter_rows(min_row=2, values_only=True)}
    assert modules == {"xor_gate", "my_xor_gate"}


def test_missing_modulefile_exits_two(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    outdir = tmp_path / "out"
    outdir.mkdir()
    argv = ["--fe_collateral", str(workarea), "--modulefile", str(tmp_path / "nope.md")]
    assert run(argv, outdir, monkeypatch) == 2
    assert list(outdir.iterdir()) == []


def test_dry_run_lists_module_workbook(
    workarea: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture
) -> None:
    modulefile = tmp_path / "modules.md"
    modulefile.write_text("xor_gate\n")
    outdir = tmp_path / "out"
    outdir.mkdir()
    argv = ["--fe_collateral", str(workarea), "--modulefile", str(modulefile), "--dry-run"]
    assert run(argv, outdir, monkeypatch) == 0
    out = capsys.readouterr().out
    assert "find_collisions.module.xlsx" in out
    assert list(outdir.iterdir()) == []

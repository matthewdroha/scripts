from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from find_collisions.collisions import RAW_FIELDS, reduced_fields
from find_collisions.report import summary_lines, write_excel_workbook, write_summary
from find_collisions.paths import resolve_inputs


def test_write_excel_workbook_creates_tables(tmp_path: Path) -> None:
    path = tmp_path / "out.xlsx"
    rows = [
        {
            "module": "m",
            "library": "LIB",
            "configrule": "parent cell's library",
            "source": '"/a.sv",1',
            "instance_count": 3,
        }
    ]
    write_excel_workbook(path, [("raw", rows, RAW_FIELDS), ("reduced", [], reduced_fields())])
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["raw", "reduced"]
    raw = workbook["raw"]
    assert [cell.value for cell in raw[1]] == RAW_FIELDS
    assert raw["E2"].value == 3
    assert list(raw.tables.values())[0].ref == "A1:E2"
    assert list(workbook["reduced"].tables.values())[0].ref == "A1:E1"


def test_write_excel_workbook_handles_configrule_columns(tmp_path: Path) -> None:
    path = tmp_path / "out.xlsx"
    fields = reduced_fields(["default library search order"])
    rows = [
        {
            "module": "m",
            "source": '"/a.sv",1',
            "library_count": 1,
            "config_rule_count": 1,
            "instance_count": 2,
            "default library search order": 1,
        }
    ]
    write_excel_workbook(path, [("reduced", rows, fields)])
    sheet = load_workbook(path)["reduced"]
    assert sheet["F1"].value == "default library search order"
    assert sheet["F2"].value == 1


def test_summary_lines_and_write(tmp_path: Path, workarea: Path) -> None:
    inputs = resolve_inputs(workarea)
    lines = summary_lines(
        ["find_collisions", "--fe_collateral", str(workarea)],
        inputs,
        [
            ("find_collisions.xlsx", [("raw", 4), ("reduced", 2)]),
            ("find_collisions.module.xlsx", [("patterns", 2), ("instances", 9)]),
        ],
        modulefile=tmp_path / "modules.md",
        notes=["instances sheet hit Excel's row limit; 5 rows dropped"],
    )
    report = tmp_path / "find_collisions.report"
    write_summary(report, lines)
    text = report.read_text()
    assert "BLOCK:     corimh" in text
    assert str(inputs.config_xml) in text
    assert "modules.md" in text
    assert "  find_collisions.xlsx\n    raw: 4\n    reduced: 2" in text
    assert "    instances: 9" in text
    assert "rows dropped" in text

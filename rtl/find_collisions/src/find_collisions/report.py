"""Report writers: the xlsx workbook and the human-readable run summary."""

from __future__ import annotations

import re
from collections.abc import Iterable, Sequence
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from find_collisions.paths import Inputs

XLSX_NAME = "find_collisions.xlsx"
MODULE_XLSX_NAME = "find_collisions.module.xlsx"
REPORT_NAME = "find_collisions.report"
EXCEL_MAX_ROWS = 1_048_576

Sheet = tuple[str, Sequence[dict[str, object]], Sequence[str]]


def _table_name(sheet_name: str) -> str:
    return "tbl_" + re.sub(r"\W", "_", sheet_name)


def write_excel_workbook(path: Path, sheets: Iterable[Sheet]) -> None:
    """Write one worksheet per (name, rows, fields), each wrapped in an Excel Table."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows, fields in sheets:
        worksheet = workbook.create_sheet(title=name)
        worksheet.append(list(fields))
        for row in rows:
            worksheet.append([row.get(field, "") for field in fields])
        ref = f"A1:{get_column_letter(len(fields))}{len(rows) + 1}"
        table = Table(displayName=_table_name(name), ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=True
        )
        worksheet.add_table(table)
    workbook.save(path)


def input_lines(inputs: Inputs, modulefile: Path | None = None) -> list[str]:
    lines = [
        f"_cfg.sv:              {inputs.cfg_sv}",
        f"rtl_list_2stage.tcl:  {inputs.rtl_list}",
        f"fullchipdump:         {inputs.fullchipdump}",
        f"config_diagnostics:   {inputs.config_xml}",
    ]
    if modulefile is not None:
        lines.append(f"modulefile:           {modulefile}")
    return lines


def summary_lines(
    argv: Sequence[str],
    inputs: Inputs,
    outputs: Sequence[tuple[str, Sequence[tuple[str, int]]]],
    modulefile: Path | None = None,
    patterns: Sequence[str] = (),
    paranoia: Sequence[str] = (),
    notes: Sequence[str] = (),
) -> list[str]:
    lines = [
        "find_collisions report",
        f"generated: {datetime.now().isoformat(timespec='seconds')}",
        f"command: {' '.join(argv)}",
        "",
        f"REF_MODEL: {inputs.ref_model}",
        f"DUT:       {inputs.dut}",
        f"BLOCK:     {inputs.block}",
        "",
        *input_lines(inputs, modulefile),
        "",
        "rows written:",
    ]
    for filename, sheet_counts in outputs:
        lines.append(f"  {filename}")
        lines += [f"    {sheet}: {count}" for sheet, count in sheet_counts]
    if patterns:
        lines.append("")
        lines.append("patterns:")
        lines += [f"  {pattern}" for pattern in patterns]
    if paranoia:
        lines.append("")
        lines += list(paranoia)
    if notes:
        lines.append("")
        lines.append("notes:")
        lines += [f"  {note}" for note in notes]
    return lines


def write_summary(path: Path, lines: Sequence[str]) -> None:
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
